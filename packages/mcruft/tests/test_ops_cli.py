# this_file: packages/mcruft/tests/test_ops_cli.py
"""Real-engine acceptance checks; all inputs and outputs live under tmp_path."""

import mcruft
import pytest
from mcruft.engines import registry
from mcruft.package import digest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pypdf import PdfReader


@pytest.fixture
def real_doc(tmp_path):
    available = {e.name for e in registry.discover() if e.available()}
    needed = {"deckcraft", "vectorcraft", "pdfcraft"}
    if missing := needed - available:
        pytest.skip(f"Missing engines: {sorted(missing)}")
    deck = tmp_path / "input.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(144), Pt(72)
    for text in ("First slide", "Second slide"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(5), Pt(5), Pt(130), Pt(50))
        box.text = text
        box.text_frame.paragraphs[0].font.name = "Arial"
        box.text_frame.paragraphs[0].font.size = Pt(12)
    prs.save(deck)
    svg = tmp_path / "input.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="96pt" height="48pt">'
        '<text x="3" y="24" font-family="Arial" font-size="12">Hello</text></svg>'
    )
    doc = mcruft.open(deck)
    doc.add_file(svg)
    return doc, deck


def test_real_pdf_exports_and_convert_policies(real_doc, tmp_path):
    doc, deck = real_doc
    for policy, convert in (
        ("editable", mcruft.convert_editable),
        ("fidelity", mcruft.convert_fidelity),
    ):
        out = tmp_path / f"document-{policy}.pdf"
        report = doc.export(out, policy=policy)
        assert report.op == "export" and report.outputs == [out]
        assert len(PdfReader(out).pages) == 3
        assert (b"/FontFile" in out.read_bytes()) == (policy == "editable")
        pdf = tmp_path / f"converted-{policy}.pdf"
        converted = convert(deck, pdf)
        assert converted.outputs == [pdf] and len(PdfReader(pdf).pages) == 2
        assert (b"/FontFile" in pdf.read_bytes()) == (policy == "editable")
    pdf = tmp_path / "converted-editable.pdf"
    for policy, convert in (
        ("editable", mcruft.convert_editable),
        ("fidelity", mcruft.convert_fidelity),
    ):
        svg = tmp_path / f"converted-{policy}.svg"
        convert(pdf, svg)
        assert ("<text" in svg.read_text()) == (policy == "editable")


def test_real_png_and_fidelity_deck(real_doc, tmp_path):
    doc, _ = real_doc
    report = doc.export(tmp_path / "pages.png", policy="fidelity", dpi=72)
    assert [p.name for p in report.outputs] == [f"pages-{i:04d}.png" for i in range(3)]
    for page, output in zip(doc.pages, report.outputs):
        with Image.open(output) as image:
            assert image.size == tuple(round(v) for v in page.size)
            assert image.convert("RGB").getextrema() != ((255, 255),) * 3
    deck = tmp_path / "fidelity.pptx"
    doc.export(deck, policy="fidelity", dpi=72)
    prs = Presentation(deck)
    assert len(prs.slides) == 3
    assert (prs.slide_width / 12700, prs.slide_height / 12700) == doc.pages[0].size
    assert all(len(slide.shapes) == 1 and slide.shapes[0].shape_type == 13 for slide in prs.slides)


def test_real_save_previews_reopen(real_doc, tmp_path):
    doc, _ = real_doc
    doc.save(tmp_path / "preview.mcruft", previews=True)
    loaded = mcruft.load(doc.path)
    assert len(loaded.pages) == 3
    for n, page in enumerate(loaded.pages):
        asset = loaded.assets[page.preview]
        assert asset.path.name == f"page-{n:04d}.pdf"
        assert asset.sha256 == digest(asset.path)
        assert asset.role == "preview" and asset.derived_from == page.source["asset"]
        assert asset.via["policy"] == "editable"
        assert len(PdfReader(asset.path).pages) == 1
        assert b"/FontFile" in asset.path.read_bytes()


def test_real_mixed_raster_document(real_doc, tmp_path):
    doc, _ = real_doc
    if not any(e.name == "photocraft" and e.available() for e in registry.discover()):
        pytest.skip("photocraft-cli not installed")
    png = tmp_path / "image.png"
    Image.new("RGB", (96, 48), "red").save(png)
    doc.add_file(png)
    for policy in ("editable", "fidelity"):
        output = tmp_path / f"mixed-{policy}.pdf"
        doc.export(output, policy=policy)
        assert len(PdfReader(output).pages) == 4
    doc.save(tmp_path / "mixed.mcruft", previews=True)
    loaded = mcruft.load(doc.path)
    assert len(loaded.pages) == 4
    for page in loaded.pages:
        asset = loaded.assets[page.preview]
        assert len(PdfReader(asset.path).pages) == 1
        assert asset.sha256 == digest(asset.path)
    report = loaded.export(tmp_path / "mixed.png", policy="fidelity", dpi=72)
    assert len(report.outputs) == 4
    for page, output in zip(loaded.pages, report.outputs):
        with Image.open(output) as image:
            assert image.size == tuple(round(v) for v in page.size)
