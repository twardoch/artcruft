# this_file: packages/mcruft/tests/test_lib.py
import pytest
from mcruft import UnsupportedRoute
from mcruft.engines.lib import engine
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader


def test_inspect_when_pdf_rotated_then_swaps_dimensions(pdf):
    assert engine.inspect(pdf) == [(200, 300), (500, 400)]


def test_inspect_when_pptx_then_returns_slide_sizes(tmp_path):
    path = tmp_path / "slides.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(path)
    assert engine.inspect(path) == [(1152, 648)] * 2


@pytest.mark.parametrize(
    "suffix,dpi,expected", [("png", (144, 144), (72, 36)), ("jpg", None, (144, 72))]
)
def test_inspect_when_image_then_uses_dpi(tmp_path, suffix, dpi, expected):
    path = tmp_path / f"image.{suffix}"
    Image.new("RGB", (144, 72)).save(path, **({"dpi": dpi} if dpi else {}))
    assert engine.inspect(path)[0] == pytest.approx(expected, rel=0.001)


@pytest.mark.parametrize(
    "attrs,expected",
    [
        ('viewBox="0 0 200 100"', (150, 75)),
        ('width="25.4mm" height="1in"', (72, 72)),
        ('width="12pt" height="16px"', (12, 12)),
    ],
)
def test_inspect_when_svg_then_returns_points(tmp_path, attrs, expected):
    path = tmp_path / "shape.svg"
    path.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" {attrs}/>')
    assert engine.inspect(path)[0] == pytest.approx(expected)


@pytest.mark.parametrize(
    "attrs,expected",
    [
        ('width="210mm" height="297mm" viewBox="0 0 210 297"', (595.275591, 841.889764)),
        ('width="21cm" height="29.7cm" viewBox="0 0 210 297"', (595.275591, 841.889764)),
        ('width="200" height="100" viewBox="0 0 20 10"', (150, 75)),
        ('viewBox="0 0 200 100"', (150, 75)),
        ('width="100%" height="100%" viewBox="0 0 200 100"', (150, 75)),
        ('width="100pt" viewBox="0 0 200 100"', (150, 75)),
    ],
)
def test_inspect_when_svg_has_viewbox_then_prefers_absolute_dimensions(tmp_path, attrs, expected):
    path = tmp_path / "sized.svg"
    path.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" {attrs}/>')
    assert engine.inspect(path)[0] == pytest.approx(expected)


def test_inspect_when_svg_size_missing_then_raises(tmp_path):
    path = tmp_path / "shape.svg"
    path.write_text("<svg/>")
    with pytest.raises(UnsupportedRoute):
        engine.inspect(path)


@pytest.mark.parametrize("policy", ["editable", "fidelity"])
def test_run_when_extract_combine_then_preserves_order(pdf, tmp_path, policy):
    one, both = tmp_path / "one.pdf", tmp_path / "both.pdf"
    report = engine.run("extract_pages", pdf, one, policy=policy, pages=[1, 0])
    assert report.outputs == [one]
    assert engine.inspect(one) == [(500, 400), (200, 300)]
    engine.run("combine", one, both, policy=policy, paths=[one, pdf])
    assert len(PdfReader(both).pages) == 4
    with pytest.raises((IndexError, ValueError)):
        engine.run("extract_pages", pdf, one, policy=policy, pages=[-1])
