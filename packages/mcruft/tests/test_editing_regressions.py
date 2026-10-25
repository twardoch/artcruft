# this_file: packages/mcruft/tests/test_editing_regressions.py
import mcruft
import pytest
from mcruft.engines import registry
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pypdf import PdfWriter

from .fakes import Fake


def test_edit_deck_copies_native_asset_and_keeps_index(isolated, tmp_path):
    path = tmp_path / "source.pptx"
    deck = Presentation()
    for text in ("First editable slide", "Second editable slide"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Pt(1), Pt(1), Pt(100), Pt(30)).text = text
    deck.save(path)
    original = path.read_bytes()
    doc = mcruft.open(path)
    page = doc.pages[1]
    with page.edit_in("deckcraft") as working:
        assert working.suffix == ".pptx" and working.read_bytes() == original
        edited = Presentation(working)
        assert len(edited.slides) == 2
        edited.slides[1].shapes[0].text = "Edited text"
        edited.save(working)
    assert page.source["index"] == 1
    assert Presentation(working).slides[1].shapes[0].text == "Edited text"
    assert path.read_bytes() == original


def test_edit_photo_copies_png_and_refreshes_size(isolated, tmp_path):
    path = tmp_path / "image.png"
    Image.new("RGB", (20, 30)).save(path)
    doc = mcruft.open(path)
    with doc.pages[0].edit_in("photocraft") as working:
        assert working.suffix == ".png" and working.read_bytes() == path.read_bytes()
        Image.new("RGB", (40, 50)).save(working)
    assert doc.pages[0].size == (40, 50)
    assert doc.pages[0].source["index"] == 0


def test_edit_pdf_refreshes_page_size(isolated, pdf):
    doc = mcruft.open(pdf)
    with doc.pages[1].edit_in("pdfcraft") as working:
        writer = PdfWriter()
        writer.add_blank_page(100, 200)
        writer.add_blank_page(300, 400)
        writer.write(working)
    assert doc.pages[1].size == (300, 400)


def test_edit_photo_rejects_fidelity_only_slide_route(isolated, tmp_path):
    path = tmp_path / "source.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    fake = Fake(("export_pages", "pptx", "pdf", "editable"), ("convert", "pdf", "psd", "fidelity"))
    registry.register(fake)
    doc = mcruft.open(path)
    with pytest.raises(mcruft.UnsupportedRoute), doc.pages[0].edit_in("photocraft"):
        pytest.fail("A fidelity-only route must never yield an editable working file")
    assert not fake.calls, "Editable handoff must not render an intermediate PDF"


def test_edit_non_native_uses_one_editable_conversion(isolated, tmp_path):
    path = tmp_path / "image.png"
    Image.new("RGB", (20, 30)).save(path)
    fake = Fake(
        ("convert", "png", "vectorcraft", "editable"),
        ("inspect", "vectorcraft", "vectorcraft", "editable"),
    )
    fake.inspect = lambda path: [(40, 50)]
    registry.register(fake)
    doc = mcruft.open(path)
    with doc.pages[0].edit_in("vectorcraft") as working:
        working.write_bytes(working.read_bytes() + b"edited")
    assert [(c[0], c[3]) for c in fake.calls] == [("convert", "editable")]
    assert doc.pages[0].size == (40, 50)


def test_edit_pdf_from_deck_uses_one_editable_conversion(isolated, tmp_path):
    source = tmp_path / "source.pptx"
    deck = Presentation()
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    fake = Fake(("convert", "pptx", "pdf", "editable"))

    def convert(op, src, dst, *, policy, **opts):
        fake.calls.append((op, policy))
        writer = PdfWriter()
        writer.add_blank_page(100, 200)
        writer.add_blank_page(300, 400)
        writer.write(dst)
        return mcruft.Report(op, policy, "fake", outputs=[dst])

    fake.run = convert
    registry.register(fake)
    doc = mcruft.open(source)
    with doc.pages[1].edit_in("pdfcraft") as working:
        assert mcruft.open(working).pages[0].size == (300, 400)
        working.write_bytes(working.read_bytes() + b"edited")
    assert fake.calls == [("convert", "editable")]
    assert doc.pages[1].source["index"] == 0
