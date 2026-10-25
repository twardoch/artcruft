# this_file: packages/mcruft_deckcraft/tests/test_deckcraft.py
import pytest
from mcruft import UnsupportedRoute
from mcruft.engines.cli import run_cli
from mcruft_deckcraft import engine
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt
from pypdf import PdfReader

pytestmark = pytest.mark.skipif(not engine.available(), reason="deckcraft-cli not installed")


def test_unknown_route(tmp_path):
    with pytest.raises(UnsupportedRoute):
        engine.run("convert", tmp_path / "in.unknown", tmp_path / "out.unknown", policy="editable")


def test_declared_routes(tmp_path):
    src = tmp_path / "in.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(144), Pt(72)
    for text in ("First", "Second"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Pt(2), Pt(2), Pt(80), Pt(40)).text = text
    prs.save(src)
    native = tmp_path / "in.deckcraft"
    run_cli([engine.binary(), "convert", src, native])
    png = tmp_path / "in.png"
    Image.new("RGB", (96, 48), "red").save(png)
    inputs = {"pptx": src, "deckcraft": native, "png": png}
    exercised = set()
    for n, cap in enumerate(engine.capabilities()):
        dst = tmp_path / f"route-{n}.{cap.dst}"
        if cap.op == "inspect":
            assert engine.inspect(inputs[cap.src]) == [(144, 72)] * 2
        else:
            report = engine.run(
                cap.op,
                inputs[cap.src],
                dst,
                policy=cap.policy,
                pages=[1, 0, 1],
                dpi=144,
                images=[png, png],
                size=(144, 72),
            )
            assert all(p.stat().st_size > 0 for p in report.outputs)
            if cap.dst == "pdf":
                pages = PdfReader(dst).pages
                assert len(pages) == (3 if cap.op == "export_pages" else 2)
                assert [p.extract_text().strip() for p in pages] == (
                    ["Second", "First", "Second"]
                    if cap.op == "export_pages"
                    else ["First", "Second"]
                )
                assert any(loss.kind == "raster" for loss in report.losses)
            if cap.dst == "png":
                assert report.outputs == [
                    dst.with_name(f"{dst.stem}-{i:04d}.png") for i in [1, 0, 1]
                ]
                for path in report.outputs:
                    with Image.open(path) as image:
                        assert image.size == (288, 144)
            if cap.dst == "pptx":
                deck = Presentation(dst)
                assert len(deck.slides) == 2
                if cap.op == "build_deck":
                    assert (deck.slide_width / 12700, deck.slide_height / 12700) == (144, 72)
                    for slide in deck.slides:
                        assert len(slide.shapes) == 1
                        shape = slide.shapes[0]
                        assert shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                        assert (shape.left, shape.top, shape.width, shape.height) == (
                            0,
                            0,
                            Pt(144),
                            Pt(72),
                        )
                    assert report.losses[0].detail == "slides are images"
            if cap.dst == "md":
                assert "First" in dst.read_text() and "Second" in dst.read_text()
        exercised.add((cap.op, cap.src, cap.dst, cap.policy))
    assert {
        (c.op, c.src, c.dst, c.policy) for c in engine.capabilities() if c.verified_with
    } <= exercised
    for ext, policy in (("pdf", "editable"), ("png", "fidelity")):
        report = engine.run("export_pages", src, tmp_path / f"all.{ext}", policy=policy)
        assert (
            len(PdfReader(report.outputs[0]).pages) == 2
            if ext == "pdf"
            else len(report.outputs) == 2
        )
    with pytest.raises(ValueError):
        engine.run("build_deck", png, tmp_path / "empty.pptx", policy="fidelity", images=[])
